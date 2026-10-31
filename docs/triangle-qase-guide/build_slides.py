"""Build an editable 16:9 PowerPoint and a matching PDF. No app code is changed."""
from pathlib import Path
import textwrap

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
A = ROOT / "assets"
OUT = ROOT.parent
W, H = 960, 540
BG, INK, MUTED = "F8FAFD", "18243B", "536279"
PURPLE, LIGHT, LINE = "6453CC", "EEEFFC", "D8DEEA"
PASS, FAIL, WHITE = "087D69", "B42338", "FFFFFF"
LAB = "https://triangle.theerapakg.moe/"
QASE = "https://app.qase.io/project/TR"
REPO = "https://github.com/tiara-stack/triangle"
WORKFLOW = REPO + "/actions/workflows/performative.yml"
RUN = "https://app.qase.io/run/TR/dashboard/16"
FONT = "/usr/share/fonts/truetype/dejavu/"
for name, file in [("Deck", "DejaVuSans.ttf"), ("Deck-Bold", "DejaVuSans-Bold.ttf"), ("Deck-Mono", "DejaVuSansMono.ttf")]:
    pdfmetrics.registerFont(TTFont(name, FONT + file))

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = "Triangle Test Lab: Qase demonstration"
prs.core_properties.subject = "Setup, four cases, manual execution and automated reporting"
prs.core_properties.author = "Triangle Test Lab"
pdf = canvas.Canvas(str(OUT / "triangle-qase-slides.pdf"), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor("Triangle Test Lab")
slide = None
n = 0


def rect(x, y, w, h, fill, stroke=None, radius=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if stroke:
        shape.line.color.rgb = RGBColor.from_string(stroke); shape.line.width = Pt(.7)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = .12
    pdf.setFillColor(HexColor("#" + fill))
    if stroke: pdf.setStrokeColor(HexColor("#" + stroke)); pdf.setLineWidth(.7)
    if radius: pdf.roundRect(x, H-y-h, w, h, radius, fill=1, stroke=bool(stroke))
    else: pdf.rect(x, H-y-h, w, h, fill=1, stroke=bool(stroke))
    return shape


def lines_for(text, width, size, font):
    lines=[]
    for original in text.split("\n"):
        if not original:
            lines.append(""); continue
        current=""
        for word in original.split():
            trial=(current + " " + word).strip()
            if current and pdfmetrics.stringWidth(trial,font,size)>width:
                lines.append(current); current=word
            else: current=trial
        lines.append(current)
    return lines


def text(txt, x, y, width, size=21, color=INK, bold=False, leading=None, mono=False, url=None):
    font="Deck-Mono" if mono else ("Deck-Bold" if bold else "Deck")
    leading=leading or size*1.24
    lines=lines_for(txt,width,size,font)
    height=len(lines)*leading+3
    if y < 0 or x < 0 or x+width>W+.2 or y+height>H+3:
        raise ValueError(f"Out-of-bounds on slide {n}: {txt[:80]} at {x},{y} {width}x{height}")
    box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(width),Pt(height))
    tf=box.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_before=p.space_after=Pt(0); p.line_spacing=Pt(leading)
        r=p.add_run(); r.text=line
        r.font.name="DejaVu Sans Mono" if mono else "DejaVu Sans"
        r.font.size=Pt(size); r.font.bold=bold
        r.font.color.rgb=RGBColor.from_string(color)
        if url: r.hyperlink.address=url
    pdf.setFillColor(HexColor("#"+color)); pdf.setFont(font,size)
    for i,line in enumerate(lines):
        pdf.drawString(x,H-y-size*.82-i*leading,line)
    if url: pdf.linkURL(url,(x,H-y-height,x+width,H-y),relative=0,thickness=0)
    return y+len(lines)*leading


def picture(file,x,y,width,max_h=None):
    path=A/file
    im=Image.open(path)
    height=width*im.height/im.width
    if max_h and height>max_h:
        width*=max_h/height; height=max_h
    slide.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(width),height=Pt(height))
    pdf.drawImage(str(path),x,H-y-height,width=width,height=height,mask="auto")
    return y+height


def step(num,title,body,x,y,width=872,size=18):
    rect(x,y+2,27,27,PURPLE,radius=5)
    text(str(num),x+8,y+5,20,13,WHITE,bold=True)
    end=text(title,x+42,y,width-42,size,INK,bold=True)
    if body: end=text(body,x+42,end+5,width-42,size-2,MUTED)
    return end+18


def start(section,title,subtitle="",dark=False,notes=""):
    global slide,n
    if n: pdf.showPage()
    n+=1
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    pdf.setFillColor(HexColor("#"+(INK if dark else BG)));pdf.rect(0,0,W,H,fill=1,stroke=0)
    color=WHITE if dark else INK
    text(f"TRIANGLE TEST LAB / {section.upper()}",44,25,840,11,"B9B2FF" if dark else PURPLE,bold=True)
    if title: text(title,44,60,872,34,color,bold=True)
    if subtitle: text(subtitle,44,111,872,16,"CDD5E4" if dark else MUTED)
    text("Qase project TR  |  28 Oct 2026",44,516,650,9,"A8B6CD" if dark else MUTED)
    text(f"{n:02d} / 10",864,516,65,9,"A8B6CD" if dark else MUTED)
    slide.notes_slide.notes_text_frame.text=notes.strip()


start("Demonstration", "", dark=True, notes="""
Purpose: demonstrate that specification-based tests expose two known classifier defects.
This is a performative testing exercise, not a performance test.
Use the four cases from commit 785557b. The lab shows only the faulty classifier output.
A failed assertion is the desired evidence for a defective case. Keep the expected value correct.
Suggested order: setup before class, register four cases, start GitHub automation, execute the manual run while it works, then inspect the two Qase reports.
Links: https://triangle.theerapakg.moe/ ; https://app.qase.io/project/TR
""")
text("Catch two defects\nwith four test cases",44,67,850,43,WHITE,bold=True,leading=51)
text("Set up Qase, run the cases manually, then report automated checks.",44,184,855,19,"CDD5E4")
picture("lab-ac-input-result.png",44,247,565)
text("2 passed",658,266,245,36,"65D7BE",bold=True)
text("2 failed",658,319,245,36,"FF93A2",bold=True)
text("Expected in the manual run\nand each automated suite.",658,378,250,17,"CDD5E4",leading=23)
text("Open lab",44,484,130,13,"B9B2FF",url=LAB)
text("Open Qase",185,484,140,13,"B9B2FF",url=QASE)
text("Open GitHub",337,484,180,13,"B9B2FF",url=REPO)

start("Setup 1 / Qase", "Start with a Qase project", "Prepare access and the place where the cases will live.",notes="""
Audience prerequisites: a browser, a Qase account or workspace invitation, permission to create cases and runs, and access to Triangle Lab.
In Qase Projects, reuse Triangle with project code TR. For an independent exercise, create a project in your workspace and note its project code. Use that same code in GitHub and .env.local.
Create the suite 'Triangle demo - four cases' from Repository > Suites > plus.
Vocabulary: Project contains the product's tests. Suite groups cases. Case defines steps, data and expected result. Run records an execution and its actual outcomes.
The screenshot is the real TR repository, with the Manual test > Create manually menu open. The case counts can change.
Source: https://docs.qase.io/en/articles/5563688-getting-started
""")
picture("qase-create-menu-crop.png",44,156,872)
y=step(1,"Sign in or accept a workspace invitation","Open Projects, then reuse Triangle, code TR, or create your own project.",44,338,size=18)
step(2,"Create a suite in Repository","Use the plus beside Suites. Name it Triangle demo - four cases.",44,y,size=18)
text("Project > Suite > Case > Run",44,474,872,20,PURPLE,bold=True)

start("Setup 2 / Reporting", "Connect the two reporters", "Configure this once in your Qase workspace and GitHub repository.",notes="""
Qase Apps: open Vitest, activate it if needed, open Access tokens, create a token. Repeat for Playwright. The current workspace already has both apps installed.
GitHub: use tiara-stack/triangle with permission, or fork the repo and enable Actions in your fork. Open repository Settings > Secrets and variables > Actions.
On Variables, add repository variable QASE_TESTOPS_PROJECT=TR, or your own Qase project code.
On Secrets, add QASE_VITEST_API_TOKEN with the Vitest app token and QASE_PLAYWRIGHT_API_TOKEN with the Playwright app token.
Alternative: a shared token can use QASE_TESTOPS_API_TOKEN. Separate app tokens are preferable for app attribution. Do not confuse the GitHub setting names with the reporter env name: scripts/performative-test.ts maps each selected token into QASE_TESTOPS_API_TOKEN for its child process.
The workflow and reporter packages are already in this repo. There is no need to install a Qase GitHub integration just to upload these results.
Check the configuration log says 'Qase TestOps reporting is enabled.' Then check four results in each new Qase run.
Never show real tokens on slides or commit them.
Sources: https://developers.qase.io/docs/prerequisites ; https://developers.qase.io/docs/playwright ; https://docs.github.com/en/actions/how-tos/write-workflows/choose-what-workflows-do/use-secrets
""")
text("Qase Apps > Vitest and Playwright > Activate > Access tokens",44,156,872,18,bold=True)
text("Create one token for each app.",44,187,872,17,MUTED)
text("GitHub > Settings > Secrets and variables > Actions",44,232,872,18,bold=True)
rect(44,270,872,34,INK)
text("Setting name",57,278,515,14,WHITE,bold=True)
text("Where / value",590,278,300,14,WHITE,bold=True)
for i,(name,value) in enumerate([
    ("QASE_TESTOPS_PROJECT","Variable / TR"),
    ("QASE_VITEST_API_TOKEN","Secret / Vitest app token"),
    ("QASE_PLAYWRIGHT_API_TOKEN","Secret / Playwright app token"),
]):
    y=304+i*45;rect(44,y,872,45,WHITE if i%2==0 else LIGHT)
    text(name,57,y+12,515,16,mono=True)
    text(value,590,y+12,310,16)
text("The existing workflow installs the reporters and sends both suites to Qase.",44,462,872,17,MUTED)

start("Qase repository", "Register each case with an expected result", "Use the specification when writing the case. Record the faulty output during execution.",notes="""
Repository > Suites plus: create 'Triangle demo - four cases'.
Manual test > Create manually opens the full case editor.
Title: Classifies a = c as Isosceles. Choose the suite. In this workspace Status is called Actual. Automation status Manual, Type Functional.
Pre-conditions: Triangle Lab open, Enter sides only selected.
Test Case Steps: Classic > New step.
Action: Enter the three side lengths, click Run, read the classifier output.
Data: a=2; b=1; c=2. Expected result: Isosceles.
Use Save and create another for the remaining cases. Use the next slide as the answer sheet.
Qase assigns real case IDs with the TR prefix. Keep them.
Screenshot: an unsaved draft in the real project. No cases were created for this document.
Source: https://docs.qase.io/en/articles/5563704-test-cases
""")
step(1,"Manual test > Create manually","Choose the suite. Set Automation status to Manual and Type to Functional.",44,157,width=872,size=18)
text("Title",44,246,120,14,MUTED,bold=True)
text("Classifies a = c as Isosceles",172,242,740,23,bold=True)
text("Classic step: action, data, expected result",44,286,872,16,PURPLE,bold=True)
picture("qase-case-steps-crop.png",44,315,872,max_h=174)
text("Repeat with the other three inputs and expected results on the next slide.",44,493,872,12,MUTED)

start("The answer sheet", "Four inputs, four expected results", "These are the four cases in the current demonstration.",notes="""
Case 01: a=2,b=1,c=2. Correct specification Isosceles. Faulty classifier Scalene. Record Failed.
Case 02: a=1,b=200,c=200. Correct specification Isosceles. Faulty classifier Value is out of range. Record Failed.
Case 03: a=3,b=3,c=3. Expected and observed Equilateral. Record Passed.
Case 04: a=201,b=1,c=1. Expected and observed Value is out of range. Record Passed.
The allowed side range is inclusive 1 through 200. The valid-200 case proves an inclusive upper boundary defect. The a=c case proves an omitted equality check.
Use 03,01,02,04 as a convenient manual demonstration order: a passing baseline, the two defects, then a passing out-of-range rejection.
Source: src/lib/triangle.ts, tests/performative/triangle.test.ts, tests/e2e/performative.spec.ts at commit 785557b.
Specification: https://triangle.theerapakg.moe/lesson-3-triangle-problem.pdf
""")
widths=[212,116,170,250,124]
xs=[44]
for w in widths:xs.append(xs[-1]+w)
rect(44,156,872,43,INK)
for i,title in enumerate(["Case","a, b, c","Expected","Actual output","Record"]):text(title,xs[i]+11,169,widths[i]-19,14,WHITE,bold=True)
rows=[
    ("01  a and c match","2, 1, 2","Isosceles","Scalene","Failed",FAIL),
    ("02  200 is allowed","1, 200, 200","Isosceles","Value is out of range","Failed",FAIL),
    ("03  Equilateral","3, 3, 3","Equilateral","Equilateral","Passed",PASS),
    ("04  201 rejected","201, 1, 1","Value is out of range","Value is out of range","Passed",PASS),
]
for j,row in enumerate(rows):
    y=199+j*62;rect(44,y,872,62,WHITE if j%2==0 else LIGHT)
    for i,value in enumerate(row[:5]):
        text(value,xs[i]+11,y+15,widths[i]-19,15.5,row[5] if i==4 else INK,bold=i in [0,4],leading=19)
text("Whole numbers 1 through 200 are allowed, including 200.",44,465,872,18,PURPLE,bold=True)

start("Manual execution", "Run, compare, record in Qase", "The expected result belongs to the case. The actual result comes from Triangle Lab.",notes="""
Select only the four manual cases in the Qase Repository and click Run for an express run. Name the run 'Triangle demo - Manual'. Alternatively use Test Runs > Start new test run and choose the four cases.
Open a case in the run or Open wizard. Keep the lab in another tab.
In Triangle Lab choose Enter sides only, enter a,b,c and click Run. Read Observed / Faulty classifier output.
Passed means actual equals the case's expected result. Failed means they differ. An HTTP 200 alone does not make the classifier correct.
In Add result, select Failed, add the Comment shown, attach the lab screenshot, and save the result. Do not change the expected result to the faulty output.
After all four cases, verify 2 Passed and 2 Failed, then Complete the manual run.
The Add result image is a clearly labeled official Qase Help Center example because no real manual run or result was created to illustrate this guide.
Sources: https://docs.qase.io/en/articles/5563702-test-runs ; image https://docs.qase.io/en/articles/6417206-github
""")
y=step(1,"Select the four manual cases > Run","Name the run Triangle demo - Manual. Open its wizard.",44,157,width=530,size=18)
y=step(2,"Enter sides in the lab > Run","Read the faulty classifier's output, then compare it with the case.",44,y,width=530,size=18)
step(3,"Save Passed or Failed in Qase","For a failure, add a comment and attach the lab screenshot.",44,y,width=530,size=18)
rect(44,418,530,73,LIGHT,radius=6)
text("Example Comment",58,430,498,13,PURPLE,bold=True)
text("a=2, b=1, c=2. Expected: Isosceles. Actual: Scalene.",58,452,498,16,leading=22)
bottom=picture("qase-add-result.png",626,156,290,max_h=317)
text("Add result dialog. Official Qase help example.",626,bottom+9,290,10,MUTED)

start("Automated execution", "Start the existing GitHub workflow", "In your repository or fork, run both suites once and inspect Qase.",notes="""
GitHub repository > Actions > Performative defect detection.
Click Run workflow, choose main, and click Run workflow again.
Workflow URL: https://github.com/tiara-stack/triangle/actions/workflows/performative.yml
The workflow installs dependencies and Chromium, reports Qase configuration status, then runs pnpm test:performative.
The script runs the direct Effect/Vitest tests and then the Playwright UI tests even when the direct suite fails.
Playwright builds and starts the app from this checkout on a local port. These checks do not depend on a running production deployment.
Inspect failures for the a=c and valid-200 assertions. Look for enabled reporting and upload messages.
A red workflow is intentional for these known defects, but setup/network/authentication failures are not evidence of a classifier defect. Open the test logs and verify the failing assertion.
GitHub requires write access to manually run a workflow; a fork can be used for an independent exercise.
Source: https://docs.github.com/en/actions/how-tos/manage-workflow-runs/manually-run-a-workflow
""")
for i,(title,body) in enumerate([("GitHub Actions","Runs the checked-out tests"),("Vitest + Playwright","Four checks in each suite"),("Qase project TR","Two reports with results")]):
    x=44+i*302
    rect(x,156,268,85,LIGHT,radius=7)
    text(title,x+15,173,240,20,bold=True)
    text(body,x+15,207,240,13,MUTED)
    if i<2:text(">",x+278,180,22,26,PURPLE,bold=True)
y=step(1,"Actions > Performative defect detection","Click Run workflow, choose main, then Run workflow again.",44,276,size=19)
step(2,"Check logs, then open Qase Test Runs","Confirm reporting is enabled and that both suites uploaded four results.",44,y,size=19)
text("Open workflow",44,457,225,17,PURPLE,bold=True,url=WORKFLOW)
text("Two assertion failures per suite should make the job red.",300,457,616,16,FAIL,bold=True)

start("Automated results", "Inspect the uploaded Qase results", "Open the newest Direct Effect and Browser UI runs. Each should have four checks.",notes="""
Actual screenshot: Qase Browser UI run 16 from GitHub workflow run 40702371101, commit 785557b.
Run 15 is Direct Effect, started by Vitest. Run 16 is Browser UI, started by Playwright.
Each has two Failed and two Passed results, so there are eight results across both runs.
Open a Failed result. The a=c browser result expected an Isosceles heading and reports that it was not found. Direct tests show the expected/observed object mismatch.
Case 01 and case 02 fail in both suites. Case 03 and case 04 pass in both suites.
Qase may show a 100% completion rate, 4 of 4 results, while run status still says In Progress. Use Complete to close the run if needed. Do not treat the lifecycle status as the test outcome.
Traceability: current code has no explicit Qase case IDs. Reporters match title and suite path. The automated repository entries may be separate from the manual cases. To share a case history, add the actual Qase IDs in both test files. Never invent IDs from this slide's 01-04 numbering.
Do not dispatch a new workflow just to view these existing example results.
Links: https://app.qase.io/run/TR/dashboard/15 ; https://app.qase.io/run/TR/dashboard/16
Source: https://developers.qase.io/docs/linking-tests
""")
text("Example: Browser UI, run 16",44,153,650,18,bold=True,url=RUN)
picture("qase-run-table.png",44,192,872)
text("2 passed + 2 failed in each suite",44,467,680,22,PURPLE,bold=True)
text("Open a Failed case for assertion details.",44,495,872,12,MUTED)

start("Optional / Local setup", "Follow along on your own machine", "Use Git, Node.js 22.18 and pnpm 11.9.0 to match the repository workflow.",notes="""
Manual testing needs only a browser and Qase. These local commands are for someone who wants to run the existing automation.
Clone https://github.com/tiara-stack/triangle.git or your fork, cd triangle, pnpm install --frozen-lockfile.
Install Chromium with node node_modules/@playwright/test/cli.js install chromium. On Linux, add --with-deps if system libraries are missing.
Copy .env.local.example to .env.local. Set QASE_TESTOPS_PROJECT, QASE_VITEST_API_TOKEN and QASE_PLAYWRIGHT_API_TOKEN. Leave shared QASE_TESTOPS_API_TOKEN empty when using distinct app tokens.
The file is already gitignored. Keep real token values private.
Run pnpm test:performative:qase to execute both suites and upload. The script loads .env.local via Effect ConfigProvider; shell environment variables take precedence. It maps each app token to its corresponding reporter process.
Run QASE_MODE=off pnpm test:performative for a rehearsal without publishing. Use pnpm test:performative:direct or pnpm test:performative:browser to select one suite. Add --qase to those scripts to publish.
Playwright builds and starts the local app automatically. A nonzero exit due to the two known assertion failures is expected.
Reporters are already installed: vitest-qase-reporter and playwright-qase-reporter. Conditional reporter configuration is in vitest.node.config.ts and playwright.config.ts. The script is scripts/performative-test.ts.
""")
text("1  Install from the repository",44,157,420,18,bold=True)
rect(44,193,424,151,LIGHT,radius=6)
text("git clone " + chr(92) + "\n  https://github.com/tiara-stack/triangle.git\ncd triangle\npnpm install --frozen-lockfile",58,209,397,12.5,mono=True,leading=24)
text("Install Chromium",44,365,423,16,bold=True)
text("node node_modules/@playwright/test/cli.js \\\n  install chromium",44,394,424,13,mono=True,leading=21)
text("2  Copy .env.local.example to .env.local",502,157,414,17,bold=True)
rect(502,193,414,190,LIGHT,radius=6)
text("QASE_TESTOPS_PROJECT=TR\n\nQASE_VITEST_API_TOKEN=<Vitest token>\n\nQASE_PLAYWRIGHT_API_TOKEN=<Playwright token>",516,208,386,12,mono=True,leading=21)
text("Replace the placeholders with your own tokens.",502,398,414,13,MUTED)
rect(44,453,872,43,INK,radius=6)
text("3  pnpm test:performative:qase",59,465,841,18,WHITE,mono=True)

start("Demo checklist", "Show the evidence, then the conclusion", "Use the passing cases to establish a baseline and the failures to demonstrate detection.",notes="""
Before presenting: verify Qase access, four manual cases, GitHub tokens/variable and workflow permission. Rehearse once and keep the lab plus Qase tabs open.
For speed, start GitHub automation first. While it works, execute the manual cases in order Equilateral, a=c, valid 200, 201 rejected.
Complete the manual run after verifying 2 Passed and 2 Failed.
Open the newest Direct Effect and Browser UI Qase runs. Each should have 4 results, 2 Passed and 2 Failed. Open one failed assertion to show the evidence.
Explain: the classifier is faulty, and correct tests expose it. The expected result must remain the specification, not the observed faulty output.
Case registration and result entry are separate actions. Creating a manual case does not automatically connect it to an automated test. The current code matches titles and suite paths, with no explicit Qase IDs. See the handout for this traceability note.
Primary references: Qase first steps https://docs.qase.io/en/articles/5563688-getting-started ; Qase cases https://docs.qase.io/en/articles/5563704-test-cases ; Qase runs https://docs.qase.io/en/articles/5563702-test-runs ; reporter tokens https://developers.qase.io/docs/prerequisites ; linking https://developers.qase.io/docs/linking-tests ; GitHub manual dispatch https://docs.github.com/en/actions/how-tos/manage-workflow-runs/manually-run-a-workflow
Image sources: real Triangle Lab and Qase project TR captures on 28 Oct 2026. Add result screenshot from official Qase Help Center GitHub article, https://docs.qase.io/en/articles/6417206-github .
The guide and deck did not create any real test case, submit a result, or dispatch a workflow.
""")
y=step(1,"Before the session","Qase access, four manual cases, GitHub settings and one rehearsal.",44,159,size=19)
y=step(2,"During the demonstration","Start automation. Run manual cases 03, 01, 02, 04. Complete the manual run.",44,y,size=19)
step(3,"Show the reports","Each run has 2 passed + 2 failed. Open a failed assertion and explain the mismatch.",44,y,size=19)
text("The classifier is faulty. These tests catch it.",44,433,872,26,PURPLE,bold=True)
text("Triangle Lab",44,482,200,14,PURPLE,url=LAB)
text("Qase repository",274,482,220,14,PURPLE,url=QASE)
text("GitHub workflow",534,482,250,14,PURPLE,url=WORKFLOW)

pdf.save()
prs.save(OUT / "triangle-qase-slides.pptx")
print(OUT / "triangle-qase-slides.pptx")
print(OUT / "triangle-qase-slides.pdf")
